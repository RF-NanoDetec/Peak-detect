"""
Generate a 2-slide Sanofi pitch deck for the Peak Analysis Tool.
Design: Light mode with orange accent — matching the app's light theme.
Requires: pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os

# ─── Palette (matches app's globals.css light theme) ────────
BG_WHITE = RGBColor(0xF8, 0xF9, 0xFA)      # --background: 210 17% 98%
CARD_WHITE = RGBColor(0xFF, 0xFF, 0xFF)     # --card: pure white
BORDER = RGBColor(0xE2, 0xE8, 0xF0)         # --border: 214 32% 91%
BORDER_LIGHT = RGBColor(0xF1, 0xF5, 0xF9)   # Very light border

ORANGE = RGBColor(0xF9, 0x73, 0x16)         # --accent: 24 95% 53%
ORANGE_SOFT = RGBColor(0xFB, 0x92, 0x3C)    # Softer orange
ORANGE_LIGHT = RGBColor(0xFF, 0xED, 0xD5)   # Very light orange tint
ORANGE_BG = RGBColor(0xFF, 0xF7, 0xED)      # Faintest orange bg

RED_DOT = RGBColor(0xE5, 0x3B, 0x3D)        # Logo red (#e53b3d)

TEXT_DARK = RGBColor(0x0F, 0x17, 0x2A)      # --foreground: 222 47% 11%
TEXT_BODY = RGBColor(0x33, 0x41, 0x55)       # Slightly lighter body
TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)      # --muted-foreground
TEXT_DIM = RGBColor(0x94, 0xA3, 0xB8)        # Dimmer text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

LOGO_PATH = "logo_light.png"

# ─── Typography ─────────────────────────────────────────────
FONT_DISPLAY = "Segoe UI Light"
FONT_HEADING = "Segoe UI Semibold"
FONT_BODY = "Segoe UI"
FONT_MONO = "Cascadia Code"


# ─── Helpers ────────────────────────────────────────────────
def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, left, top, width, height, fill_color, border_color=None, border_w=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_w or 1)
    else:
        shape.line.fill.background()
    return shape


def rrect(slide, left, top, width, height, fill_color, border_color=None, border_w=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_w)
    else:
        shape.line.fill.background()
    return shape


def textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def set_text(tf, text, size=14, color=TEXT_DARK, bold=False, align=PP_ALIGN.LEFT, font=FONT_BODY):
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return p


def accent_line(slide, left, top, width, color=ORANGE, thickness=3):
    return rect(slide, left, top, width, Pt(thickness), color)


def add_logo(slide, left, top, height):
    """Add the logo PNG to the slide."""
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, left, top, height=height)


def build_footer(slide):
    """Shared footer for both slides — light theme."""
    # Footer background
    rect(slide, Inches(0), Inches(6.88), SLIDE_WIDTH, Inches(0.62), WHITE)
    # Orange top-line on footer
    accent_line(slide, Inches(0), Inches(6.87), SLIDE_WIDTH, ORANGE, thickness=2.5)

    # Left: attribution
    tb = textbox(slide, Inches(0.8), Inches(6.97), Inches(6), Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT

    r1 = p.add_run()
    r1.text = "Dr. Lucjan Grzegorzewski"
    r1.font.size = Pt(10)
    r1.font.color.rgb = TEXT_DARK
    r1.font.bold = True
    r1.font.name = FONT_HEADING

    r2 = p.add_run()
    r2.text = "    Hamburg University"
    r2.font.size = Pt(10)
    r2.font.color.rgb = TEXT_MUTED
    r2.font.name = FONT_BODY

    # Right: Sanofi
    tb2 = textbox(slide, Inches(9.5), Inches(6.97), Inches(3.3), Inches(0.4))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r3 = p2.add_run()
    r3.text = "Sanofi  |  Technical Overview"
    r3.font.size = Pt(9)
    r3.font.color.rgb = TEXT_DIM
    r3.font.name = FONT_BODY


def feature_card(slide, left, top, width, height, number, title, desc):
    """Feature card with left orange accent stripe on white card."""
    card = rrect(slide, left, top, width, height, CARD_WHITE, border_color=BORDER, border_w=1)

    # Left orange stripe
    rect(slide, left + Inches(0.02), top + Inches(0.06), Pt(4), height - Inches(0.12), ORANGE)

    # Number
    tb_num = textbox(slide, left + Inches(0.2), top + Inches(0.04), Inches(0.42), Inches(0.3))
    p = tb_num.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"0{number}"
    r.font.size = Pt(16)
    r.font.color.rgb = ORANGE
    r.font.bold = True
    r.font.name = FONT_MONO

    # Title
    tb_t = textbox(slide, left + Inches(0.62), top + Inches(0.06), width - Inches(0.8), Inches(0.26))
    set_text(tb_t.text_frame, title, size=11, color=TEXT_DARK, bold=True, font=FONT_HEADING)

    # Description
    tb_d = textbox(slide, left + Inches(0.62), top + Inches(0.32), width - Inches(0.8), Inches(0.4))
    set_text(tb_d.text_frame, desc, size=8.5, color=TEXT_MUTED, font=FONT_BODY)


def advantage_card(slide, left, top, width, height, icon_text, title, desc):
    """Performance advantage card with orange-tinted background."""
    card = rrect(slide, left, top, width, height, ORANGE_LIGHT, border_color=ORANGE_SOFT, border_w=1)

    # Icon/symbol
    tb_icon = textbox(slide, left + Inches(0.12), top + Inches(0.08), Inches(0.35), Inches(0.3))
    p = tb_icon.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = icon_text
    r.font.size = Pt(16)
    r.font.color.rgb = ORANGE
    r.font.bold = True
    r.font.name = FONT_BODY

    # Title
    tb_t = textbox(slide, left + Inches(0.48), top + Inches(0.06), width - Inches(0.6), Inches(0.24))
    set_text(tb_t.text_frame, title, size=10, color=TEXT_DARK, bold=True, font=FONT_HEADING)

    # Description
    tb_d = textbox(slide, left + Inches(0.48), top + Inches(0.3), width - Inches(0.6), Inches(0.35))
    set_text(tb_d.text_frame, desc, size=8, color=TEXT_BODY, font=FONT_BODY)


# ────────────────────────────────────────────────────────────
# SLIDE 1 — Overview
# ────────────────────────────────────────────────────────────
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, BG_WHITE)

    # Top accent bar
    accent_line(slide, Inches(0), Inches(0), SLIDE_WIDTH, ORANGE, thickness=4)

    # ── Logo top-right ──
    add_logo(slide, Inches(10.8), Inches(0.3), Inches(1.1))

    # ── Title block ──
    tb1 = textbox(slide, Inches(0.8), Inches(0.4), Inches(9), Inches(0.65))
    tf1 = tb1.text_frame
    tf1.word_wrap = False
    p = tf1.paragraphs[0]

    r1 = p.add_run()
    r1.text = "PEAK "
    r1.font.size = Pt(36)
    r1.font.color.rgb = ORANGE
    r1.font.bold = True
    r1.font.name = FONT_DISPLAY

    r2 = p.add_run()
    r2.text = "ANALYSIS TOOL"
    r2.font.size = Pt(36)
    r2.font.color.rgb = TEXT_DARK
    r2.font.bold = False
    r2.font.name = FONT_DISPLAY

    # Tagline
    tb2 = textbox(slide, Inches(0.8), Inches(1.0), Inches(9), Inches(0.35))
    set_text(tb2.text_frame,
             "Interactive time-trace analysis for single-particle detection, counting & characterization",
             size=13, color=TEXT_MUTED, font=FONT_BODY)

    accent_line(slide, Inches(0.8), Inches(1.42), Inches(3.0), ORANGE, thickness=3)

    # ──────────────────────────────────────────────────────
    # ROW 1: CAPABILITIES — "What you can do"
    # ──────────────────────────────────────────────────────
    section_y = Inches(1.6)
    tb_sec1 = textbox(slide, Inches(0.8), section_y, Inches(3), Inches(0.28))
    set_text(tb_sec1.text_frame, "CAPABILITIES", size=10, color=ORANGE, bold=True, font=FONT_HEADING)

    capabilities = [
        ("Fast data loading",
         "Dead-time correction, parallelized preprocessing for files with 20M+ data points"),
        ("Signal processing",
         "Butterworth & Savitzky-Golay filtering for noise reduction and baseline correction"),
        ("Peak detection & analysis",
         "Peak width, height, prominence extraction with parameter histograms"),
        ("Time trace analysis",
         "Drift detection and temporal trends across measurement runs"),
        ("Double peak analysis",
         "Aggregate & doublet identification with scatter-plot metrics"),
        ("Data export",
         "Peak property tables (CSV, Excel) and publication-ready PNG images"),
    ]

    card_w = Inches(3.85)
    card_h = Inches(0.7)
    col_x = [Inches(0.8), Inches(4.78), Inches(8.76)]
    row1_y = section_y + Inches(0.32)
    row2_y = row1_y + card_h + Inches(0.1)

    for i, (title, desc) in enumerate(capabilities):
        col = i % 3
        row = i // 3
        x = col_x[col]
        y = row1_y if row == 0 else row2_y
        feature_card(slide, x, y, card_w, card_h, i + 1, title, desc)

    # ──────────────────────────────────────────────────────
    # ROW 2: PERFORMANCE — "Why it's great"
    # ──────────────────────────────────────────────────────
    perf_section_y = row2_y + card_h + Inches(0.18)
    tb_sec2 = textbox(slide, Inches(0.8), perf_section_y, Inches(3), Inches(0.28))
    set_text(tb_sec2.text_frame, "PERFORMANCE", size=10, color=ORANGE, bold=True, font=FONT_HEADING)

    advantages = [
        ("\u26A1", "Near-C speed",
         "Numba JIT-compiled core routines for fast analysis of large datasets"),
        ("\u2B24", "20M+ data points",
         "Min-max decimation preserves peak shapes while keeping rendering smooth"),
        ("\u2922", "Interactive exploration",
         "Pan, zoom & inspect millions of points with visible-range-only rendering via uPlot"),
    ]

    adv_w = Inches(3.85)
    adv_h = Inches(0.62)
    adv_y = perf_section_y + Inches(0.32)

    for i, (icon, title, desc) in enumerate(advantages):
        x = col_x[i]
        advantage_card(slide, x, adv_y, adv_w, adv_h, icon, title, desc)

    # ──────────────────────────────────────────────────────
    # BOTTOM ROW: Use Cases (left) + Technology (right)
    # ──────────────────────────────────────────────────────
    bottom_y = adv_y + adv_h + Inches(0.18)
    bottom_h = Inches(1.18)

    # ── Use Cases card (left, wider) ──
    uc_w = Inches(7.78)
    uc_card = rrect(slide, Inches(0.8), bottom_y, uc_w, bottom_h,
                    CARD_WHITE, border_color=BORDER, border_w=1)

    tb_uc = textbox(slide, Inches(1.05), bottom_y + Inches(0.08), Inches(2), Inches(0.24))
    set_text(tb_uc.text_frame, "APPLICATION", size=10, color=ORANGE, bold=True, font=FONT_HEADING)
    accent_line(slide, Inches(1.05), bottom_y + Inches(0.32), Inches(0.8), ORANGE, thickness=2)

    # Application description — paragraph style
    tb_app = textbox(slide, Inches(1.05), bottom_y + Inches(0.4), uc_w - Inches(0.5), Inches(0.7))
    tf_app = tb_app.text_frame
    tf_app.word_wrap = True
    p = tf_app.paragraphs[0]
    p.space_before = Pt(2)

    segments = [
        ("Single-particle detection & counting", True),
        (" in nanochannels by fluorescence or scattering.  ", False),
        ("Characterization", True),
        (" of throughput, velocity, and peak parameters (width, height, area).  ", False),
        ("Multi-channel analysis", True),
        (" for different wavelengths ", False),
        ("(coming soon)", False),
        (".", False),
    ]
    for text, bold in segments:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(9.5)
        r.font.color.rgb = TEXT_BODY if not bold else TEXT_DARK
        r.font.bold = bold
        r.font.name = FONT_BODY

    # ── Technology card (right) ──
    tech_x = Inches(8.76)
    tech_w = Inches(3.85)
    ts_card = rrect(slide, tech_x, bottom_y, tech_w, bottom_h,
                    CARD_WHITE, border_color=BORDER, border_w=1)

    tb_ts = textbox(slide, tech_x + Inches(0.2), bottom_y + Inches(0.08), Inches(2), Inches(0.24))
    set_text(tb_ts.text_frame, "TECHNOLOGY", size=10, color=ORANGE, bold=True, font=FONT_HEADING)
    accent_line(slide, tech_x + Inches(0.2), bottom_y + Inches(0.32), Inches(0.8), ORANGE, thickness=2)

    techs = [
        ("Backend", "FastAPI  NumPy  SciPy  Numba JIT"),
        ("Frontend", "React  TypeScript  uPlot"),
        ("Tooling", "uv (Rust)  PyInstaller  WebSocket"),
    ]
    tb_tsl = textbox(slide, tech_x + Inches(0.2), bottom_y + Inches(0.42), tech_w - Inches(0.4), Inches(0.7))
    tf_tsl = tb_tsl.text_frame
    tf_tsl.word_wrap = True
    for i, (label, stack) in enumerate(techs):
        p = tf_tsl.paragraphs[0] if i == 0 else tf_tsl.add_paragraph()
        p.space_before = Pt(5)
        r_label = p.add_run()
        r_label.text = f"{label}   "
        r_label.font.size = Pt(8.5)
        r_label.font.color.rgb = ORANGE
        r_label.font.bold = True
        r_label.font.name = FONT_MONO
        r_stack = p.add_run()
        r_stack.text = stack
        r_stack.font.size = Pt(8.5)
        r_stack.font.color.rgb = TEXT_MUTED
        r_stack.font.name = FONT_MONO

    # ── Footer ──
    build_footer(slide)


# ────────────────────────────────────────────────────────────
# SLIDE 2 — Live Demo
# ────────────────────────────────────────────────────────────
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, BG_WHITE)

    # Top accent bar
    accent_line(slide, Inches(0), Inches(0), SLIDE_WIDTH, ORANGE, thickness=4)

    # ── Logo top-right ──
    add_logo(slide, Inches(10.8), Inches(0.3), Inches(1.1))

    # ── Title ──
    tb1 = textbox(slide, Inches(0.8), Inches(0.45), Inches(10), Inches(0.7))
    tf1 = tb1.text_frame
    tf1.word_wrap = False
    p = tf1.paragraphs[0]
    r1 = p.add_run()
    r1.text = "LIVE "
    r1.font.size = Pt(38)
    r1.font.color.rgb = ORANGE
    r1.font.bold = True
    r1.font.name = FONT_DISPLAY
    r2 = p.add_run()
    r2.text = "DEMO"
    r2.font.size = Pt(38)
    r2.font.color.rgb = TEXT_DARK
    r2.font.bold = False
    r2.font.name = FONT_DISPLAY

    tb2 = textbox(slide, Inches(0.8), Inches(1.2), Inches(10), Inches(0.35))
    set_text(tb2.text_frame, "Complete analysis workflow in ~30 seconds",
             size=14, color=TEXT_MUTED, font=FONT_BODY)

    accent_line(slide, Inches(0.8), Inches(1.65), Inches(2.5), ORANGE, thickness=3)

    # ── Embedded video ──
    frame_l = Inches(0.8)
    frame_t = Inches(1.95)
    frame_w = Inches(11.7)
    frame_h = Inches(4.0)

    VIDEO_PATH = "Video Project 1.mp4"

    # Placeholder frame — video will be injected post-save
    frame = rrect(slide, frame_l, frame_t, frame_w, frame_h,
                  CARD_WHITE, border_color=BORDER, border_w=2)

    cx = frame_l + frame_w // 2
    cy = frame_t + frame_h // 2

    circle_size = Inches(0.9)
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        cx - circle_size // 2, cy - circle_size // 2,
        circle_size, circle_size
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = ORANGE
    circle.line.fill.background()

    tri_size = Inches(0.35)
    tri = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        cx - tri_size // 2 + Inches(0.04), cy - tri_size // 2,
        tri_size, tri_size
    )
    tri.rotation = 90.0
    tri.fill.solid()
    tri.fill.fore_color.rgb = WHITE
    tri.line.fill.background()

    tb3 = textbox(slide, Inches(3.5), cy + Inches(0.6), Inches(6.3), Inches(0.4))
    set_text(tb3.text_frame, "Click to play demo video",
             size=12, color=TEXT_DIM, align=PP_ALIGN.CENTER, font=FONT_BODY)

    # ── Workflow pipeline ──
    steps = ["Load", "Preprocess", "Detect", "Analyze", "Export"]

    total_steps = len(steps)
    pill_w = Inches(1.85)
    pill_h = Inches(0.52)
    gap = Inches(0.35)
    total_w = total_steps * pill_w + (total_steps - 1) * gap
    start_x = (SLIDE_WIDTH - total_w) // 2
    step_y = Inches(6.2)

    for i, step in enumerate(steps):
        x = start_x + i * (pill_w + gap)

        # Pill background (white card with orange accent)
        pill = rrect(slide, x, step_y, pill_w, pill_h,
                     fill_color=CARD_WHITE, border_color=ORANGE, border_w=1)

        # Step text
        tb_s = textbox(slide, x, step_y, pill_w, pill_h)
        tf_s = tb_s.text_frame
        tf_s.word_wrap = False
        p = tf_s.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(7)

        # Number
        rn = p.add_run()
        rn.text = f"{i + 1}  "
        rn.font.size = Pt(13)
        rn.font.color.rgb = ORANGE
        rn.font.bold = True
        rn.font.name = FONT_MONO

        # Label
        rs = p.add_run()
        rs.text = step
        rs.font.size = Pt(12)
        rs.font.color.rgb = TEXT_DARK
        rs.font.name = FONT_HEADING

        # Connector arrow between pills
        if i < total_steps - 1:
            arr_x = x + pill_w
            arr_tb = textbox(slide, arr_x, step_y, gap, pill_h)
            pf = arr_tb.text_frame
            pf.paragraphs[0].alignment = PP_ALIGN.CENTER
            pf.paragraphs[0].space_before = Pt(5)
            ra = pf.paragraphs[0].add_run()
            ra.text = "\u2192"
            ra.font.size = Pt(14)
            ra.font.color.rgb = ORANGE_SOFT
            ra.font.name = FONT_BODY

    # ── Footer ──
    build_footer(slide)


# ────────────────────────────────────────────────────────────
# Video injection — post-process the PPTX zip to embed mp4
# ────────────────────────────────────────────────────────────
def inject_video(pptx_path, video_path, slide_index=1):
    """
    Inject an MP4 video into slide 2 of the PPTX.
    Works by manipulating the PPTX zip and XML directly.
    slide_index is 0-based (slide 2 = index 1).
    """
    import zipfile
    import shutil
    import tempfile
    from copy import deepcopy

    if not os.path.exists(video_path):
        print(f"  Video not found: {video_path}")
        return

    video_size = os.path.getsize(video_path) / (1024 * 1024)
    print(f"  Injecting video ({video_size:.1f} MB)...")

    tmp_dir = tempfile.mkdtemp()
    tmp_pptx = os.path.join(tmp_dir, "output.pptx")

    slide_num = slide_index + 1  # 1-based for file paths
    slide_xml_path = f"ppt/slides/slide{slide_num}.xml"
    slide_rels_path = f"ppt/slides/_rels/slide{slide_num}.xml.rels"
    video_part_path = "ppt/media/video1.mp4"
    content_types_path = "[Content_Types].xml"

    with zipfile.ZipFile(pptx_path, 'r') as zin:
        with zipfile.ZipFile(tmp_pptx, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)

                if item.filename == content_types_path:
                    # Add mp4 content type if missing
                    tree = etree.fromstring(data)
                    ns = tree.nsmap.get(None, '')
                    # Check if mp4 extension already registered
                    has_mp4 = any(
                        el.get('Extension') == 'mp4'
                        for el in tree.findall(f'{{{ns}}}Default')
                    )
                    if not has_mp4:
                        ext_el = etree.SubElement(tree, f'{{{ns}}}Default')
                        ext_el.set('Extension', 'mp4')
                        ext_el.set('ContentType', 'video/mp4')
                    data = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)

                elif item.filename == slide_rels_path:
                    # Add video relationship
                    tree = etree.fromstring(data)
                    ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
                    # Find highest rId
                    max_id = 0
                    for rel in tree:
                        rid = rel.get('Id', 'rId0')
                        num = int(rid.replace('rId', ''))
                        max_id = max(max_id, num)

                    # Add video relationship
                    video_rid = f'rId{max_id + 1}'
                    media_rid = f'rId{max_id + 2}'

                    rel_video = etree.SubElement(tree, f'{{{ns}}}Relationship')
                    rel_video.set('Id', video_rid)
                    rel_video.set('Type', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/video')
                    rel_video.set('Target', '../media/video1.mp4')

                    rel_media = etree.SubElement(tree, f'{{{ns}}}Relationship')
                    rel_media.set('Id', media_rid)
                    rel_media.set('Type', 'http://schemas.microsoft.com/office/2007/relationships/media')
                    rel_media.set('Target', '../media/video1.mp4')

                    data = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)

                    # Now modify the slide XML to add the video shape
                    slide_data = zin.read(slide_xml_path)
                    slide_tree = etree.fromstring(slide_data)

                    # Define namespaces
                    nsmap = {
                        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
                        'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
                        'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main',
                    }

                    # Build video shape XML
                    spTree = slide_tree.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}spTree')
                    if spTree is None:
                        spTree = slide_tree.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cSld/{http://schemas.openxmlformats.org/presentationml/2006/main}spTree')

                    # Find the shape tree
                    cSld = slide_tree.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld')
                    spTree = cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spTree')

                    # Count existing shapes for unique ID
                    shape_count = len(spTree) + 1

                    # Video dimensions matching the frame
                    left = int(Inches(0.8))
                    top = int(Inches(1.95))
                    width = int(Inches(11.7))
                    height = int(Inches(4.0))

                    # Build the pic element for video
                    a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
                    p_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
                    r_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
                    p14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'

                    pic = etree.SubElement(spTree, f'{{{p_ns}}}pic')

                    # nvPicPr
                    nvPicPr = etree.SubElement(pic, f'{{{p_ns}}}nvPicPr')
                    cNvPr = etree.SubElement(nvPicPr, f'{{{p_ns}}}cNvPr')
                    cNvPr.set('id', str(shape_count + 100))
                    cNvPr.set('name', 'Demo Video')

                    # Add a:hlinkClick for video link
                    hlinkClick = etree.SubElement(cNvPr, f'{{{a}}}hlinkClick')
                    hlinkClick.set(f'{{{r_ns}}}id', '')
                    hlinkClick.set('action', 'ppaction://media')

                    cNvPicPr = etree.SubElement(nvPicPr, f'{{{p_ns}}}cNvPicPr')
                    picLocks = etree.SubElement(cNvPicPr, f'{{{a}}}picLocks')
                    picLocks.set('noChangeAspect', '1')

                    nvPr = etree.SubElement(nvPicPr, f'{{{p_ns}}}nvPr')
                    videoFile = etree.SubElement(nvPr, f'{{{a}}}videoFile')
                    videoFile.set(f'{{{r_ns}}}link', video_rid)

                    extLst = etree.SubElement(nvPr, f'{{{p_ns}}}extLst')
                    ext = etree.SubElement(extLst, f'{{{p_ns}}}ext')
                    ext.set('uri', '{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}')
                    p14media = etree.SubElement(ext, f'{{{p14}}}media')
                    p14media.set(f'{{{r_ns}}}embed', media_rid)

                    # blipFill (blank — no poster image)
                    blipFill = etree.SubElement(pic, f'{{{p_ns}}}blipFill')
                    blip = etree.SubElement(blipFill, f'{{{a}}}blip')
                    stretch = etree.SubElement(blipFill, f'{{{a}}}stretch')
                    fillRect = etree.SubElement(stretch, f'{{{a}}}fillRect')

                    # spPr (position and size)
                    spPr = etree.SubElement(pic, f'{{{p_ns}}}spPr')
                    xfrm = etree.SubElement(spPr, f'{{{a}}}xfrm')
                    off = etree.SubElement(xfrm, f'{{{a}}}off')
                    off.set('x', str(left))
                    off.set('y', str(top))
                    ext_el = etree.SubElement(xfrm, f'{{{a}}}ext')
                    ext_el.set('cx', str(width))
                    ext_el.set('cy', str(height))
                    prstGeom = etree.SubElement(spPr, f'{{{a}}}prstGeom')
                    prstGeom.set('prst', 'rect')

                    # Write modified slide XML
                    modified_slide_data = etree.tostring(slide_tree, xml_declaration=True, encoding='UTF-8', standalone=True)
                    zout.writestr(slide_xml_path, modified_slide_data)

                if item.filename != slide_xml_path:  # slide XML written above
                    zout.writestr(item, data)

            # Add the video file
            zout.write(video_path, video_part_path)

    # Replace original with modified
    shutil.move(tmp_pptx, pptx_path)
    shutil.rmtree(tmp_dir)
    print(f"  Video embedded successfully")


# ────────────────────────────────────────────────────────────
# Main
# ────────────────────────────────────────────────────────────
def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    build_slide_1(prs)
    build_slide_2(prs)

    out = "PeakAnalysis_Sanofi_Pitch.pptx"
    prs.save(out)
    print(f"Created: {out}")

    # Post-process: inject video into slide 2
    inject_video(out, "Video Project 1.mp4", slide_index=1)


if __name__ == "__main__":
    main()
